from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "docs" / "presentation" / "Horizon_Auto_Insurance_Solution_Deck.pptx"

W, H = 20, 11.25
LIGHT = "F9F7F4"
WHITE = "FFFFFF"
NAVY = "0B2026"
INK = "1B3139"
MUTED = "5A5755"
LINE = "D8D5D0"
PALE = "EEEDE9"
LAVA = "FF3621"
LAVA_LIGHT = "FFE2DE"
TEAL = "1B5162"
GREEN = "00875C"


def rgb(value):
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def textbox(slide, text, x, y, w, h, size=20, color=INK, bold=False,
            font="Arial", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
            margin=0, tracking=None):
    font = {"DM Sans": "Arial", "DM Mono": "Courier New"}.get(font, font)
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    if tracking is not None:
        run.font._element.set("spc", str(tracking))
    return shape


def rich_text(slide, runs, x, y, w, h, size=20, color=INK, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    for value, opts in runs:
        run = p.add_run()
        run.text = value
        run.font.name = {"DM Sans": "Arial", "DM Mono": "Courier New"}.get(
            opts.get("font", "Arial"), opts.get("font", "Arial")
        )
        run.font.size = Pt(opts.get("size", size))
        run.font.bold = opts.get("bold", False)
        run.font.color.rgb = rgb(opts.get("color", color))
    return shape


def label(slide, text, x=0.8, y=0.55, w=12):
    return textbox(slide, text.upper(), x, y, w, 0.35, 13, MUTED, True, "Courier New")


def title(slide, heading, sub=None, label_text=None):
    if label_text:
        label(slide, label_text)
    textbox(slide, heading, 0.8, 1.1 if label_text else 0.65, 18.2, 0.8, 31, NAVY, True)
    if sub:
        textbox(slide, sub, 0.8, 1.95 if label_text else 1.5, 18.0, 0.65, 18, MUTED)


def footer(slide, number, section="HORIZON AUTO INSURANCE"):
    textbox(slide, section, 0.8, 10.72, 6.5, 0.25, 10, MUTED, True, "DM Mono")
    textbox(slide, f"{number:02d}", 18.55, 10.72, 0.6, 0.25, 10, MUTED, True, "DM Mono", PP_ALIGN.RIGHT)


def new_slide(prs, bg=LIGHT):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(bg)
    return slide


def card(slide, x, y, w, h, heading, body, number=None, accent=False):
    rect(slide, x, y, w, h, WHITE, LINE, True)
    if number:
        textbox(slide, number, x + 0.35, y + 0.32, 0.65, 0.38, 13, LAVA if accent else TEAL, True, "DM Mono")
        hx = x + 1.05
    else:
        hx = x + 0.38
    heading_size = 20 if h >= 2.4 else 18
    if len(body) <= 55 and h >= 2.4:
        body_size = 20
    elif len(body) <= 95 and h >= 2.4:
        body_size = 18
    elif len(body) <= 85:
        body_size = 16
    else:
        body_size = 15
    textbox(slide, heading, hx, y + 0.27, w - (hx - x) - 0.3, 0.52, heading_size, NAVY, True)
    textbox(slide, body, x + 0.38, y + 0.95, w - 0.76, h - 1.12, body_size, MUTED)


def metric(slide, x, y, w, value, heading, body="", accent=False):
    rect(slide, x, y, w, 2.05, WHITE, LINE, True)
    textbox(slide, heading.upper(), x + 0.32, y + 0.28, w - 0.64, 0.3, 11, MUTED, True, "DM Mono")
    textbox(slide, value, x + 0.32, y + 0.64, w - 0.64, 0.65, 30, LAVA if accent else NAVY, True)
    if body:
        textbox(slide, body, x + 0.32, y + 1.38, w - 0.64, 0.4, 12, MUTED)


def add_picture(slide, path, x, y, w, h=None):
    kwargs = {"left": Inches(x), "top": Inches(y), "width": Inches(w)}
    if h is not None:
        kwargs["height"] = Inches(h)
    return slide.shapes.add_picture(str(path), **kwargs)


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "Horizon Auto Insurance — proactive risk prevention and renewal decisioning"
    prs.core_properties.subject = "Business value, solution flow, and Databricks technical architecture"
    prs.core_properties.author = "Databricks"

    # 1 — title
    s = new_slide(prs, NAVY)
    textbox(s, "HORIZON AUTO INSURANCE · SOLUTION AND DEMO", 0.9, 0.75, 11, 0.4, 13, "C4CCD6", True, "DM Mono")
    rich_text(s, [("Catch driving risk ", {"color": WHITE, "bold": True, "size": 42}),
                  ("before", {"color": LAVA, "bold": True, "size": 42}),
                  (" it becomes a claim.", {"color": WHITE, "bold": True, "size": 42})],
              0.9, 2.1, 17.7, 1.25)
    textbox(s, "Proactive prevention and human-approved renewal decisions on one governed platform.",
            0.9, 3.65, 14.5, 0.8, 22, "EEEDE9")
    rect(s, 0.9, 8.55, 5.8, 0.05, LAVA)
    textbox(s, "BUSINESS → WORKFLOW → TECHNICAL DEEP DIVE", 0.9, 8.85, 8, 0.35, 12, "C4CCD6", True, "DM Mono")
    textbox(s, "Horizon Auto Insurance", 0.9, 9.45, 8, 0.45, 17, WHITE, True)
    textbox(s, "Databricks solution narrative", 0.9, 9.95, 8, 0.35, 14, "C4CCD6")

    # 2 — problem
    s = new_slide(prs)
    title(s, "Risk changes between portfolio reviews.", "Horizon often learns about worsening behavior after the loss signal becomes a claim.", "The business problem")
    rect(s, 0.8, 2.9, 11.5, 5.55, WHITE, LINE, True)
    textbox(s, "“", 1.2, 3.15, 0.6, 0.8, 46, LAVA, True)
    textbox(s, "My underwriters can explain yesterday’s book. They cannot see which drivers became riskier this week — or intervene before renewal.",
            1.85, 3.42, 9.6, 1.75, 26, NAVY, True)
    textbox(s, "Maya Chen · Regional Underwriting Manager", 1.85, 5.45, 7, 0.4, 14, MUTED, True)
    card(s, 12.75, 2.9, 6.45, 1.65, "Late signal", "Periodic review reveals unsafe trends after exposure has accumulated.", "01", True)
    card(s, 12.75, 4.85, 6.45, 1.65, "Fragmented evidence", "Telemetry, claims, guidelines, and policy state sit in separate workflows.", "02")
    card(s, 12.75, 6.8, 6.45, 1.65, "Hard to defend", "Manual recommendations lack a consistent explanation and audit trail.", "03")
    footer(s, 2)

    # 3 — stakes
    s = new_slide(prs)
    title(s, "The cost of waiting is operational and financial.", "Prototype targets shown here are modeled outcomes, not realized customer results.", "The quantified stakes")
    metric(s, 0.8, 3.0, 4.35, "40%", "Less review time", "Prioritize the policies that need judgment.", True)
    metric(s, 5.48, 3.0, 4.35, "80%", "Intervened in one day", "Reach high-risk policies while action can still matter.")
    metric(s, 10.16, 3.0, 4.35, "−5%", "Preventable claims", "Modeled reduction through earlier intervention.")
    metric(s, 14.84, 3.0, 4.35, "+2 pts", "Loss-ratio improvement", "Modeled portfolio impact from better timing.")
    rect(s, 0.8, 5.65, 18.39, 2.75, NAVY, NAVY, True)
    textbox(s, "THE DECISION WINDOW", 1.25, 6.05, 3.2, 0.35, 12, "C4CCD6", True, "DM Mono")
    textbox(s, "Unsafe behavior rises", 1.25, 6.75, 3.3, 0.45, 19, WHITE, True)
    textbox(s, "→", 4.65, 6.68, 0.6, 0.5, 24, LAVA, True)
    textbox(s, "Risk is scored", 5.35, 6.75, 2.5, 0.45, 19, WHITE, True)
    textbox(s, "→", 8.1, 6.68, 0.6, 0.5, 24, LAVA, True)
    textbox(s, "Underwriter intervenes", 8.85, 6.75, 3.7, 0.45, 19, WHITE, True)
    textbox(s, "→", 12.78, 6.68, 0.6, 0.5, 24, LAVA, True)
    textbox(s, "Renewal is decided", 13.55, 6.75, 3.2, 0.45, 19, WHITE, True)
    textbox(s, "The platform compresses the path from signal to governed action.", 1.25, 7.55, 11, 0.4, 15, "EEEDE9")
    footer(s, 3)

    # 4 — outcome
    s = new_slide(prs)
    title(s, "Detect → explain → recommend → decide → audit.", "The solution changes the underwriter’s workflow without automating the final decision.", "The outcome")
    for i, (h, b) in enumerate([
        ("Detect", "Hourly telemetry updates expose worsening driving patterns."),
        ("Explain", "The model and Genie surface the factors and governed evidence."),
        ("Recommend", "Rules translate risk into an intervention and renewal action."),
        ("Decide", "The underwriter approves or overrides with a required reason."),
        ("Audit", "Every question, recommendation, and decision is recorded."),
    ], 1):
        x = 0.8 + (i - 1) * 3.72
        card(s, x, 3.1, 3.35, 3.65, h, b, f"{i:02d}", i == 1)
        if i < 5:
            textbox(s, "→", x + 3.39, 4.45, 0.28, 0.45, 18, LAVA, True)
    rect(s, 0.8, 7.45, 18.39, 1.1, PALE, PALE, True)
    textbox(s, "CONTROL POINT", 1.15, 7.78, 2.0, 0.3, 11, LAVA, True, "DM Mono")
    textbox(s, "The model scores risk. Deterministic rules recommend. A human underwriter decides.", 3.2, 7.72, 14.8, 0.42, 17, NAVY, True)
    footer(s, 4)

    # 5 — value
    s = new_slide(prs)
    title(s, "Four levers connect the workflow to business value.", None, "Business value")
    cols = [0.8, 4.22, 9.6, 14.05, 19.2]
    headers = ["LEVER", "WHAT CHANGES", "MODELED OUTCOME", "CONTROL"]
    for x, h in zip(cols[:-1], headers):
        textbox(s, h, x + 0.18, 2.45, (cols[headers.index(h)+1] - x) - 0.3, 0.35, 11, MUTED, True, "DM Mono")
    rows = [
        ("Loss", "Intervene before renewal", "5% fewer preventable claims", "Human approval"),
        ("Speed", "One prioritized evidence view", "40% less review time", "Same governed facts"),
        ("Retention", "Reward sustained safe behavior", "3% modeled lift", "Approved program terms"),
        ("Risk", "Explain and record every action", "100% audited decisions", "Identity-aware access"),
    ]
    for r, row in enumerate(rows):
        y = 3.02 + r * 1.45
        rect(s, 0.8, y, 18.39, 1.12, WHITE, LINE, True)
        for c, value in enumerate(row):
            textbox(s, value, cols[c] + 0.18, y + 0.3, cols[c+1] - cols[c] - 0.35, 0.55, 16 if c else 18, NAVY if c == 0 else MUTED, c == 0)
    footer(s, 5)

    # 6 — why Databricks
    s = new_slide(prs)
    title(s, "One platform keeps data, intelligence, action, and audit aligned.", "No separate analytics copy, app database workflow, or AI control plane to reconcile.", "Why Databricks")
    card(s, 0.8, 3.0, 5.75, 4.3, "One governed definition", "Delta tables, a metric view, Genie, and the application resolve against the same policy-risk contract.", "01", True)
    card(s, 7.12, 3.0, 5.75, 4.3, "Operational speed", "Lakebase sync serves policy queues and decisions in milliseconds while Delta remains the analytical source of truth.", "02")
    card(s, 13.44, 3.0, 5.75, 4.3, "Built-in controls", "Unity Catalog, AI Gateway, MLflow, lineage, and app identity make access and model use observable.", "03")
    textbox(s, "The architecture is simpler because governance is shared, not copied.", 0.8, 8.15, 16, 0.55, 22, NAVY, True)
    footer(s, 6)

    # 7 — journey
    s = new_slide(prs)
    title(s, "Three milestones turn data into a governed decision.", None, "Solution journey")
    milestones = [
        ("①", "GOVERNED DATA", "Ingest policy and telematics files\nBuild trusted features\nScore claim risk", TEAL),
        ("②", "UNDERWRITER ACTION", "Prioritize the review queue\nInvestigate with Genie\nApprove or override", LAVA),
        ("③", "AUDITABLE CONTROL", "Apply user-level entitlements\nTrace model and agent calls\nPersist decisions", NAVY),
    ]
    for i, (n, h, b, color) in enumerate(milestones):
        x = 0.8 + i * 6.22
        rect(s, x, 2.7, 5.72, 5.3, WHITE, LINE, True)
        textbox(s, n, x + 0.4, 3.08, 0.8, 0.7, 28, color, True)
        textbox(s, h, x + 1.3, 3.2, 3.95, 0.35, 13, MUTED, True, "DM Mono")
        textbox(s, b, x + 0.4, 4.2, 4.85, 2.2, 20, NAVY, True)
        textbox(s, ["Lakeflow · Delta · MLflow", "Databricks App · Genie · Lakebase", "Unity Catalog · AI Gateway · Lakebase"][i], x + 0.4, 7.1, 4.85, 0.45, 12, color, True, "DM Mono")
        if i < 2:
            textbox(s, "→", x + 5.78, 4.95, 0.35, 0.5, 22, LAVA, True)
    footer(s, 7)

    # 8 — demo
    s = new_slide(prs, NAVY)
    textbox(s, "LIVE DEMO", 0.9, 0.75, 5, 0.4, 13, "C4CCD6", True, "DM Mono")
    textbox(s, "Watch one policy move from signal to decision.", 0.9, 1.45, 17.8, 0.8, 32, WHITE, True)
    steps = [
        ("01", "See the queue", "Open the West-region review queue and portfolio KPIs."),
        ("02", "Inspect evidence", "Open the policy, risk factors, trend, and expected loss."),
        ("03", "Ask why", "Genie queries governed data and cites underwriting guidelines."),
        ("04", "Make the call", "Approve the recommendation or override with a reason."),
        ("05", "Prove it", "Decision history and audit events update immediately."),
    ]
    for i, (n, h, b) in enumerate(steps):
        x = 0.9 + i * 3.75
        textbox(s, n, x, 3.25, 0.65, 0.45, 14, LAVA, True, "DM Mono")
        textbox(s, h, x, 4.0, 3.2, 0.55, 20, WHITE, True)
        textbox(s, b, x, 4.85, 3.15, 1.65, 15, "C4CCD6")
        if i < 4:
            textbox(s, "→", x + 3.25, 4.05, 0.35, 0.4, 18, LAVA, True)
    textbox(s, "Demo question", 0.9, 8.15, 2.0, 0.3, 11, "C4CCD6", True, "DM Mono")
    textbox(s, "“Why is policy POL-004217 becoming riskier, and what should we do before renewal?”", 0.9, 8.65, 17.2, 0.7, 22, WHITE, True)
    textbox(s, "08", 18.55, 10.72, 0.6, 0.25, 10, "C4CCD6", True, "DM Mono", PP_ALIGN.RIGHT)

    # 9 — architecture image
    s = new_slide(prs)
    title(s, "One decision contract connects analytics and operations.", "The app, metric view, and Genie show the same risk band, action, and expected loss.", "Technical architecture")
    add_picture(s, ROOT / "docs" / "diagrams" / "horizon-data-flow.png", 1.05, 2.65, 17.9, 7.1)
    footer(s, 9)

    # 10 — pipeline
    s = new_slide(prs)
    title(s, "Incremental by design, reconciled by contract.", "Lakeflow handles continuous files and keyed corrections without rebuilding the entire portfolio.", "Data engineering deep dive")
    add_picture(s, ROOT / "docs" / "diagrams" / "horizon-pipeline-design.png", 0.95, 2.55, 11.9, 6.85)
    card(s, 13.2, 2.65, 5.95, 1.55, "Bronze", "Auto Loader preserves lineage, manifests, rescued data, and event time.", "01", True)
    card(s, 13.2, 4.5, 5.95, 1.55, "Silver", "Auto CDC applies snapshots and keyed corrections with data-quality expectations.", "02")
    card(s, 13.2, 6.35, 5.95, 1.55, "Gold", "Materialized views produce stable policy features and review candidates.", "03")
    card(s, 13.2, 8.2, 5.95, 1.25, "After pipeline", "Score → decide → sync to Lakebase.", "04")
    footer(s, 10)

    # 11 — intelligence
    s = new_slide(prs)
    title(s, "Separate prediction, business rules, and human judgment.", "Each layer has a distinct purpose, owner, and audit boundary.", "ML and decisioning deep dive")
    layers = [
        ("MODEL", "Predict claim risk", "Interpretable features\nProbability and drivers\nMLflow registry + aliases"),
        ("RULES", "Recommend an action", "Risk band thresholds\nExpected loss\nApproved intervention logic"),
        ("HUMAN", "Record the decision", "Approve or override\nRequired rationale\nReviewer identity + timestamp"),
    ]
    for i, (k, h, b) in enumerate(layers):
        x = 0.8 + i * 6.22
        rect(s, x, 3.0, 5.72, 4.9, WHITE, LINE, True)
        textbox(s, k, x + 0.4, 3.35, 1.2, 0.32, 12, LAVA if i == 0 else MUTED, True, "DM Mono")
        textbox(s, h, x + 0.4, 4.0, 4.9, 0.55, 21, NAVY, True)
        textbox(s, b, x + 0.4, 5.0, 4.7, 1.8, 17, MUTED)
        if i < 2:
            textbox(s, "→", x + 5.78, 5.0, 0.35, 0.45, 20, LAVA, True)
    rect(s, 0.8, 8.45, 18.39, 0.8, PALE, PALE, True)
    textbox(s, "No autonomous premium change. The platform supports a governed underwriting decision.", 1.2, 8.68, 17.0, 0.35, 16, NAVY, True)
    footer(s, 11)

    # 12 — runtime
    s = new_slide(prs)
    title(s, "Fast reads, governed investigation, human-only writes.", "The signed-in user’s identity travels through the application and analytics layers.", "Application deep dive")
    add_picture(s, ROOT / "docs" / "diagrams" / "horizon-runtime-interaction.png", 0.95, 2.55, 12.4, 7.25)
    card(s, 13.7, 2.75, 5.45, 1.65, "Review", "Lakebase serves queues, evidence, and history in milliseconds.", "01", True)
    card(s, 13.7, 4.75, 5.45, 1.65, "Investigate", "Genie runs governed SQL and searches approved guidelines with citations.", "02")
    card(s, 13.7, 6.75, 5.45, 1.65, "Decide", "Only the underwriter writes a decision; every action emits an audit event.", "03")
    footer(s, 12)

    # 13 — governance
    s = new_slide(prs)
    title(s, "Governance follows the user from table to decision.", None, "Governance and responsible use")
    controls = [
        ("Access", "Unity Catalog ABAC masks customer names and filters policies by assigned region."),
        ("Lineage", "Raw telemetry, features, scores, recommendations, and metrics remain traceable."),
        ("AI", "Unity AI Gateway governs model access; Genie cites SQL and guideline sources."),
        ("Audit", "Lakebase stores review state, decisions, overrides, and question history."),
        ("Quality", "Pipeline expectations fail on orphaned keys and quarantine malformed events."),
        ("Fairness", "The model excludes direct PII and supports monitored comparison-group analysis."),
    ]
    for i, (h, b) in enumerate(controls):
        x = 0.8 + (i % 3) * 6.22
        y = 2.75 + (i // 3) * 3.0
        card(s, x, y, 5.72, 2.45, h, b, f"{i+1:02d}", i == 0)
    footer(s, 13)

    # 14 — technology map
    s = new_slide(prs)
    title(s, "Each Databricks service has one clear job.", None, "Technology map")
    rows = [
        ("Lakeflow", "Ingest files, apply CDC, enforce quality, build features"),
        ("Delta Lake + Unity Catalog", "Store governed data, permissions, lineage, and model assets"),
        ("MLflow", "Track training, register the model, manage production aliases"),
        ("Metric views + Genie", "Define shared KPIs and answer governed natural-language questions"),
        ("Lakebase", "Serve low-latency app reads and persist operational decisions"),
        ("Databricks Apps", "Deliver the underwriter workflow with user identity"),
        ("Unity AI Gateway", "Control and observe model endpoint access"),
        ("Declarative Automation Bundles", "Deploy jobs, pipelines, warehouse, app, and permissions as code"),
    ]
    for i, (tech, job) in enumerate(rows):
        col, row = i % 2, i // 2
        x = 0.8 + col * 9.35
        y = 2.55 + row * 1.65
        rect(s, x, y, 8.85, 1.25, WHITE, LINE, True)
        textbox(s, tech, x + 0.35, y + 0.25, 3.1, 0.65, 17, NAVY, True)
        textbox(s, job, x + 3.35, y + 0.25, 5.1, 0.7, 14, MUTED)
    footer(s, 14)

    # 15 — rollout
    s = new_slide(prs)
    title(s, "Pilot the workflow before scaling the book.", "Use the synthetic prototype to validate operations, controls, and value measurement.", "Pilot plan")
    phases = [
        ("01", "Validate", "One region\nOne renewal cohort\nShadow decisions", "2 weeks"),
        ("02", "Pilot", "Named underwriters\nDaily review queue\nWeekly control review", "4–6 weeks"),
        ("03", "Scale", "More regions\nChampion/challenger model\nOutcome feedback loop", "After gates"),
    ]
    for i, (n, h, b, t) in enumerate(phases):
        x = 0.8 + i * 6.22
        rect(s, x, 2.9, 5.72, 4.6, WHITE, LINE, True)
        textbox(s, n, x + 0.38, 3.25, 0.7, 0.4, 14, LAVA if i == 0 else TEAL, True, "DM Mono")
        textbox(s, h, x + 1.15, 3.2, 3.6, 0.5, 21, NAVY, True)
        textbox(s, b, x + 0.38, 4.25, 4.7, 1.7, 18, MUTED)
        textbox(s, t.upper(), x + 0.38, 6.7, 2.5, 0.32, 11, TEAL, True, "DM Mono")
    textbox(s, "Exit gates", 0.8, 8.15, 2, 0.35, 13, MUTED, True, "DM Mono")
    textbox(s, "Data quality holds · recommendations are explainable · entitlements match policy · underwriters trust the workflow · outcome metrics are measurable", 2.7, 8.08, 16.0, 0.75, 17, NAVY, True)
    footer(s, 15)

    # 16 — ask
    s = new_slide(prs, NAVY)
    textbox(s, "THE ASK", 0.9, 0.75, 3, 0.4, 13, "C4CCD6", True, "DM Mono")
    textbox(s, "Put earlier risk signals in the hands of underwriters.", 0.9, 1.55, 17.7, 0.9, 34, WHITE, True)
    card(s, 0.9, 3.2, 5.65, 3.2, "Choose the cohort", "Select one region and renewal window where earlier intervention can be observed.", "01", True)
    card(s, 7.15, 3.2, 5.65, 3.2, "Agree on controls", "Confirm rating boundaries, program terms, entitlements, and audit requirements.", "02")
    card(s, 13.4, 3.2, 5.65, 3.2, "Measure the pilot", "Track review time, intervention coverage, overrides, retention, and modeled loss impact.", "03")
    textbox(s, "One governed platform. One shared decision contract. One human accountable for the outcome.", 0.9, 8.35, 17.5, 0.65, 22, WHITE, True)
    textbox(s, "Horizon Auto Insurance · Databricks solution narrative", 0.9, 10.1, 9, 0.3, 12, "C4CCD6", True, "DM Mono")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
